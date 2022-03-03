# excelを扱うモジュール
import xlsxwriter
import openpyxl


class Excelopration:       
    def chart_style(self,chart,x_name,y_name):
        # グラフの規定スタイルの設定
        chart.set_chartarea({
            'border': {'color': 'white'}
        })
        chart.set_x_axis({
            # x軸
            'name': str(x_name), # ラベル名設定
            'num_font': {'name': 'Arial', 'size': 7,}, # 数値のフォント設定
            'name_font':{'name': 'Arial', 'size': 10, 'bold': False, 'italic': True}, # ラベルのフォント設定
            'crossing' : 0, # 軸の交点設定
            'major_tick_mark' : 'inside', # 主目盛の向き設定
            'line': {'color': "black", 'width': 0} # 軸の色と線の太さ設定
        })
        chart.set_y_axis({
            # y軸
            'name' : str(y_name[0]), 
            'num_font': {'name': 'Arial', 'size': 7},
            'name_font': {'name': 'Arial', 'size':10, 'bold': False, 'italic': True},
            'crossing' : 0, 
            'major_gridlines': {'visible': False}, 
            'major_tick_mark' : 'inside', 
            'line': {'color': 'black', 'width': 0}
        })
        chart.set_x2_axis({
            # 第二のx軸
            'num_font': {'name': 'Arial', 'size': 7,}, 
            'name_font':{'name': 'Arial', 'size': 10, 'bold': False, 'italic': True}, 
            'major_tick_mark' : 'inside', 
            'line': {'color': "black", 'width': 0} ,
            'label_position':'none',
            'crossing': 'max',
            'visible': True,
        })

        chart.set_y2_axis({
            # 第二のy軸
            'name' : str(y_name[-1]), 
            'num_font': {'name': 'Arial', 'size': 7,}, 
            'name_font':{'name': 'Arial', 'size': 10, 'bold': False, 'italic': True}, 
            'crossing' : 'max', 
            'major_gridlines': {'visible': False}, 
            'major_tick_mark' : 'inside', 
            'line': {'color': 'black', 'width': 0}
        })
   
    def excel_series(self,chart,worksheet,x_num,num,trend,y2_axis=False):
        # 取得データ選択、プロットスタイル等の設定
        chart.add_series({
            'name': [worksheet.name,0,num],
            'categories': [worksheet.name,1,0,x_num,0], # x軸の値
            'values': [worksheet.name,1,num,x_num,num], # y軸の値
            'marker': {'type':'circle'},
            'y2_axis': y2_axis,# 二軸の設定
            
            # 近似タイプや近似スタイル設定
            'trendline': {
                'type': trend,
                'display_equation': True,
                'line': {
                    'width': 1,
                    'dash_type': 'dot',
                }
            }
        })

    def excel(self,excelfile,outfile,trend):
        # データ読み込みからグラフ作成
        wb = openpyxl.load_workbook(excelfile) # 既存ファイルからデータ読み込み
        workbook = xlsxwriter.Workbook(outfile) # 保存先ファイルの作成   

        for sheet in wb:
            # ワークシートごとに処理
            worksheet = workbook.add_worksheet() 
            y_num = len(list(sheet.columns)) # 作成するグラフの列の数を取得
            x_num = len(list(sheet.rows)) # 作成するグラフの行の数を取得
            chart = workbook.add_chart({'type': 'scatter'})
            trend = trend # 近似タイプ
            x_name = ''   # x軸のカラム名
            y_name = []   # y軸のカラム名

            for num,rows in enumerate(sheet.columns):
                # カラムごとに処理
                if num==0:
                    for n,cell in enumerate(rows):
                        if n==0:
                            x_name = cell.value
                        worksheet.write(n,0,cell.value) # 保存先に値を記入
                else:
                    for n,cell in enumerate(rows):
                        if n==0:
                            y_name.append(cell.value)
                        worksheet.write(n,num,cell.value)

                    if num==1:
                        self.excel_series(chart,worksheet,x_num,num,trend,y2_axis=False)                        
                    else:
                        self.excel_series(chart,worksheet,x_num,num,trend,y2_axis=True)
                    
                    self.chart_style(chart,x_name,y_name)
                    worksheet.insert_chart('F4',chart) # グラフの貼り付け先

        workbook.close()
